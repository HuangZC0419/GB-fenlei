# -*- coding: utf-8 -*-
"""把批量分类结果导出为带格式的 Excel 工作簿。"""

import unicodedata
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import CATEGORIES, DEFAULT_SYSTEM_MAP, DEFAULT_THRESHOLD, OUTPUT_DIR


class ExportError(Exception):
    """可预期的导出错误，携带稳定错误码。"""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(size=10)
THIN_SIDE = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

CONFIDENCE_BINS = ("高（≥0.80）", "中（0.60–0.79）", "低（<0.60）")

MANUAL_CONFIRMED = "人工"
MANUAL_PENDING = "待人工"
RESULT_PROFESSION_HEADER = "所属专业名称"
RESULT_CONFIDENCE_HEADER = "置信度"

DEFAULT_RESULT_HEADERS = [
    "序号",
    "标准号",
    "标准名称",
    "选用系统名称",
    "选用专业名称",
    "置信度",
    "人工复核",
]


def result_headers(payload: Mapping[str, Any]) -> List[str]:
    """结果表列名沿用上传参照表的原始用词，避免自行改名。"""
    columns = payload.get("columns") or {}
    headers = list(DEFAULT_RESULT_HEADERS)
    headers[1] = _text(columns.get("code")) or headers[1]
    headers[2] = _text(columns.get("name")) or headers[2]
    headers[3] = _text(columns.get("system")) or headers[3]
    headers[4] = _text(columns.get("profession")) or headers[4]
    return headers


def classification_filename(directory: Optional[Path] = None, now: Optional[datetime] = None) -> str:
    """生成带时间戳的导出文件名；同一秒内重复导出自动加序号。"""
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    base = "分类结果_{}".format(stamp)
    if directory is None:
        return base + ".xlsx"
    folder = Path(directory)
    candidate = base + ".xlsx"
    counter = 2
    while (folder / candidate).exists():
        candidate = "{}_{}.xlsx".format(base, counter)
        counter += 1
    return candidate


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _confidence(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _display_width(value: Any) -> int:
    text = _text(value)
    width = 0
    for char in text:
        width += 2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
    return width


def _safe_value(value: Any) -> Any:
    """避免以等号开头的文本被 Excel 当作公式执行。"""
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def _set_header(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGNMENT
    cell.border = BORDER


def _set_body(cell, value: Any, number_format: Optional[str] = None) -> None:
    cell.value = _safe_value(value)
    cell.font = BODY_FONT
    cell.alignment = BODY_ALIGNMENT
    cell.border = BORDER
    if number_format:
        cell.number_format = number_format


def _style_range(worksheet: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in worksheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            cell.border = BORDER


def _write_table(
    worksheet: Worksheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    start_row: int = 1,
    widths: Optional[Mapping[int, float]] = None,
    number_formats: Optional[Mapping[int, str]] = None,
    freeze: bool = True,
    autofilter: bool = True,
) -> int:
    """写入一张表，返回下一张表可用的起始行。"""
    formats = number_formats or {}
    for column, header in enumerate(headers, start=1):
        _set_header(worksheet.cell(row=start_row, column=column, value=header))

    current = start_row + 1
    for record in rows:
        for column, value in enumerate(record, start=1):
            _set_body(
                worksheet.cell(row=current, column=column),
                value,
                formats.get(column),
            )
        current += 1

    last_row = max(start_row, current - 1)
    for column in range(1, len(headers) + 1):
        letter = get_column_letter(column)
        if widths and column in widths:
            width = float(widths[column])
        else:
            longest = _display_width(headers[column - 1])
            for record in rows:
                if column <= len(record):
                    longest = max(longest, _display_width(record[column - 1]))
            width = min(max(longest + 2, 10), 55)
        worksheet.column_dimensions[letter].width = width

    if autofilter:
        worksheet.auto_filter.ref = "A{0}:{1}{2}".format(
            start_row, get_column_letter(len(headers)), last_row
        )
    if freeze:
        worksheet.freeze_panes = "A{}".format(start_row + 1)
    worksheet.sheet_view.showGridLines = False
    return current


def _system_for(profession: str, mapping: Mapping[str, Any]) -> str:
    profession_to_system = mapping.get("profession_to_system") or {}
    return _text(profession_to_system.get(profession)) or DEFAULT_SYSTEM_MAP.get(profession, "")


def _normalize_overrides(
    overrides: Optional[Mapping[Any, Any]],
    row_indices: Sequence[int],
    categories: Sequence[str],
) -> Dict[int, str]:
    """校验并规范化人工复核结果：行序号 -> 专业名称。"""
    valid_indices = set(row_indices)
    valid_categories = set(categories)
    normalized: Dict[int, str] = {}
    for key, value in (overrides or {}).items():
        try:
            index = int(key)
        except (TypeError, ValueError):
            raise ExportError("INVALID_OVERRIDE", "人工复核的序号无效：{}".format(key))
        if index not in valid_indices:
            raise ExportError("INVALID_OVERRIDE", "人工复核的序号不存在：{}".format(index))
        profession = _text(value)
        if profession not in valid_categories:
            raise ExportError("INVALID_OVERRIDE", "未知的专业名称：{}".format(profession or "空"))
        normalized[index] = profession
    return normalized


def _prepare_rows(
    payload: Mapping[str, Any], overrides: Mapping[int, str]
) -> List[Dict[str, Any]]:
    mapping = payload.get("mapping") or {}
    prepared: List[Dict[str, Any]] = []
    for position, row in enumerate(payload.get("rows") or [], start=1):
        try:
            index = int(row.get("index", position))
        except (TypeError, ValueError):
            index = position
        predicted = _text(row.get("profession"))
        override = overrides.get(index)
        profession = override or predicted
        system = _text(row.get("system"))
        if override and profession != predicted:
            mapped = _system_for(profession, mapping)
            if mapped:
                system = mapped
        prepared.append({
            "index": index,
            "code": _text(row.get("code")),
            "name": _text(row.get("name")),
            "system": system,
            "profession": profession,
            "confidence": _confidence(row.get("confidence")),
            "override": override or "",
        })
    return prepared


def _result_values(row: Mapping[str, Any]) -> List[Any]:
    return [
        row["index"],
        row["code"],
        row["name"],
        row["system"],
        row["profession"],
        row["confidence"],
        row["override"],
    ]


def _confidence_bin(confidence: float) -> str:
    if confidence >= 0.80:
        return CONFIDENCE_BINS[0]
    if confidence >= 0.60:
        return CONFIDENCE_BINS[1]
    return CONFIDENCE_BINS[2]


def _statistics_rows(
    rows: Sequence[Mapping[str, Any]], payload: Mapping[str, Any], threshold: float
) -> List[List[Any]]:
    summary = payload.get("summary") or {}
    reviewed = sum(1 for row in rows if row["override"])
    pending_low = sum(
        1 for row in rows if row["confidence"] < threshold and not row["override"]
    )
    result: List[List[Any]] = [
        ["总览", "原始行数", summary.get("raw_rows", len(rows)), None],
        ["总览", "去重后标准数", summary.get("unique_rows", len(rows)), None],
        ["总览", "低置信度（未复核）", pending_low, None],
        ["总览", "人工复核条数", reviewed, None],
    ]

    dimensions: List[Tuple[str, Counter]] = [
        ("专业", Counter(row["profession"] for row in rows)),
        ("系统", Counter(row["system"] for row in rows)),
        ("置信度区间", Counter(_confidence_bin(row["confidence"]) for row in rows)),
    ]
    for dimension, counter in dimensions:
        total = sum(counter.values())
        for value, count in sorted(counter.items(), key=lambda item: (-item[1], item[0])):
            result.append([dimension, value, count, float(count) / total if total else 0.0])
        result.append([dimension, "合计", total, 1.0 if total else 0.0])
    return result


def _conflict_rows(payload: Mapping[str, Any]) -> List[List[Any]]:
    mapping = payload.get("mapping") or {}
    rows: List[List[Any]] = []

    profession_to_system = mapping.get("profession_to_system") or {}
    system_to_profession = mapping.get("system_to_profession") or {}
    rows.append([
        "映射概览",
        "专业与系统一对一" if mapping.get("is_one_to_one") else "专业与系统非一对一",
        "",
        "",
        len(profession_to_system),
        "专业 {} 个，系统 {} 个".format(
            len(profession_to_system), len(system_to_profession)
        ),
        "参照表统计",
    ])
    for profession in sorted(profession_to_system):
        rows.append([
            "映射关系",
            "专业→系统",
            profession,
            "",
            1,
            "对应系统：{}".format(profession_to_system[profession]),
            "参照表多数票",
        ])

    for issue in mapping.get("issues") or []:
        target = _text(issue.get("profession")) or _text(issue.get("system"))
        values = issue.get("values") or []
        rows.append([
            "映射问题",
            _text(issue.get("type")),
            target,
            "",
            "",
            "；".join(_text(value) for value in values),
            "需人工确认",
        ])

    for conflict in payload.get("conflicts") or []:
        labels = conflict.get("labels") or []
        chosen = _text(conflict.get("chosen"))
        rows.append([
            "冲突",
            _text(conflict.get("type")),
            _text(conflict.get("code")),
            _text(conflict.get("name")),
            conflict.get("count", ""),
            "；".join(_text(label) for label in labels),
            "采用多数票：{}".format(chosen) if chosen else "需人工确认",
        ])

    for report in payload.get("dedup_reports") or []:
        rows.append([
            "待分类去重",
            _text(report.get("type")),
            _text(report.get("code")),
            _text(report.get("name")),
            report.get("count", ""),
            _text(report.get("detail")),
            _text(report.get("action")),
        ])

    for warning in payload.get("warnings") or []:
        rows.append(["提醒", "参照表校验", "", "", "", _text(warning), "相关记录已跳过或忽略"])
    return rows


def export_pending_workbook(
    pending_path: Path,
    payload: Mapping[str, Any],
    overrides: Optional[Mapping[Any, Any]] = None,
    output_path: Optional[Path] = None,
    categories: Optional[Sequence[str]] = None,
) -> Path:
    """在原始待分类 Excel 后追加「所属专业名称」「置信度」两列。

    原始工作表、列和行全部保留；同一标准的多条重复行、多单位行都会回填相同结果。
    自动分类行写模型置信度；人工分类行写“人工”或“待人工”。
    """
    source = Path(pending_path)
    if not source.is_file():
        raise ExportError("FILE_NOT_FOUND", "待分类原文件不存在：{}".format(source.name))

    if output_path is None:
        folder = Path(OUTPUT_DIR)
        output_path = folder / classification_filename(folder)
    output_path = Path(output_path)

    meta = payload.get("pending") or {}
    sheet_name = _text(meta.get("sheet"))
    try:
        workbook = openpyxl.load_workbook(source)
    except Exception as exc:  # noqa: BLE001
        raise ExportError(
            "EXPORT_FAILED",
            "无法打开待分类原文件：{}（{}）".format(source.name, exc),
        )

    try:
        worksheet = (
            workbook[sheet_name]
            if sheet_name and sheet_name in workbook.sheetnames
            else workbook.worksheets[0]
        )
        try:
            header_row = int(meta.get("header_row") or 1)
        except (TypeError, ValueError):
            header_row = 1
        header_row = min(max(1, header_row), max(1, worksheet.max_row))

        rows = list(payload.get("rows") or [])
        row_indices = [
            int(row.get("index", position))
            for position, row in enumerate(rows, start=1)
        ]
        valid_categories = list(categories or [
            item["profession"]
            for item in payload.get("categories") or []
            if item.get("profession")
        ] or CATEGORIES)
        normalized_overrides = _normalize_overrides(
            overrides, row_indices, valid_categories
        )

        row_values: Dict[int, Tuple[str, Any]] = {}
        for row in rows:
            index = int(row.get("index", 0))
            override = normalized_overrides.get(index)
            status = _text(row.get("status")) or "auto"
            if override:
                profession = override
                confidence: Any = MANUAL_CONFIRMED
            elif status == "manual":
                profession = ""
                confidence = MANUAL_PENDING
            else:
                profession = _text(row.get("profession"))
                confidence = _confidence(row.get("confidence"))
            for source_row in row.get("source_rows") or []:
                try:
                    row_values[int(source_row)] = (profession, confidence)
                except (TypeError, ValueError):
                    continue

        last_column = max(1, worksheet.max_column)
        profession_column = last_column + 1
        confidence_column = last_column + 2
        template = worksheet.cell(row=header_row, column=last_column)

        for column, title in (
            (profession_column, RESULT_PROFESSION_HEADER),
            (confidence_column, RESULT_CONFIDENCE_HEADER),
        ):
            cell = worksheet.cell(row=header_row, column=column, value=title)
            cell.font = copy(template.font)
            cell.fill = copy(template.fill)
            cell.border = copy(template.border)
            cell.alignment = copy(template.alignment)

        for row_number in range(header_row + 1, worksheet.max_row + 1):
            if row_number not in row_values:
                continue
            profession, confidence = row_values[row_number]
            profession_cell = worksheet.cell(
                row=row_number, column=profession_column, value=profession or None
            )
            confidence_cell = worksheet.cell(
                row=row_number, column=confidence_column, value=confidence
            )
            profession_cell.alignment = copy(
                worksheet.cell(row=row_number, column=last_column).alignment
            )
            confidence_cell.alignment = copy(
                worksheet.cell(row=row_number, column=last_column).alignment
            )
            if isinstance(confidence, (int, float)):
                confidence_cell.number_format = "0.0%"

        worksheet.column_dimensions[
            get_column_letter(profession_column)
        ].width = 22
        worksheet.column_dimensions[
            get_column_letter(confidence_column)
        ].width = 12

        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_path)
        except PermissionError:
            raise ExportError(
                "FILE_LOCKED",
                "文件正被占用，无法写入：{}。请关闭 Excel 后重试。".format(
                    output_path.name
                ),
            )
        except OSError as exc:
            raise ExportError(
                "EXPORT_FAILED",
                "写入 Excel 失败：{}（{}）".format(output_path.name, exc),
            )
    finally:
        workbook.close()

    return output_path


def export_results(
    payload: Mapping[str, Any],
    overrides: Optional[Mapping[Any, Any]] = None,
    output_path: Optional[Path] = None,
    threshold: Optional[float] = None,
    categories: Optional[Sequence[str]] = None,
) -> Path:
    """将分类结果写入 Excel，返回实际文件路径。

    结果列名沿用上传参照表的原始用词；低置信度工作表在导出时按阈值重新计算，
    已人工复核的行视为处理完毕，不再列入。
    """
    if output_path is None:
        folder = Path(OUTPUT_DIR)
        output_path = folder / classification_filename(folder)
    output_path = Path(output_path)

    raw_rows = list(payload.get("rows") or [])
    row_indices: List[int] = []
    for position, row in enumerate(raw_rows, start=1):
        try:
            row_indices.append(int(row.get("index", position)))
        except (TypeError, ValueError):
            row_indices.append(position)
    valid_categories = list(categories or [
        item["profession"]
        for item in payload.get("categories") or []
        if item.get("profession")
    ] or CATEGORIES)
    normalized_overrides = _normalize_overrides(
        overrides, row_indices, valid_categories
    )
    prepared = _prepare_rows(payload, normalized_overrides)
    headers = result_headers(payload)

    summary = payload.get("summary") or {}
    if threshold is None:
        threshold = _confidence(summary.get("threshold", DEFAULT_THRESHOLD))
    threshold = _confidence(threshold)

    column_widths = {1: 8, 2: 20, 3: 40, 4: 24, 5: 20, 6: 12, 7: 16}

    workbook = openpyxl.Workbook()
    try:
        result_sheet = workbook.active
        result_sheet.title = "分类结果"
        _write_table(
            result_sheet,
            headers,
            [_result_values(row) for row in prepared],
            number_formats={6: "0.0000"},
            widths=column_widths,
        )

        low_sheet = workbook.create_sheet("低置信度")
        low_rows = [
            row
            for row in prepared
            if row["confidence"] < threshold and not row["override"]
        ]
        _write_table(
            low_sheet,
            headers,
            [_result_values(row) for row in low_rows],
            number_formats={6: "0.0000"},
            widths=column_widths,
        )

        conflict_sheet = workbook.create_sheet("数据体检")
        _write_table(
            conflict_sheet,
            ["类别", "类型", "标准号/对象", "标准名称", "数量", "详情", "处理"],
            _conflict_rows(payload),
            widths={1: 14, 2: 24, 3: 24, 4: 30, 5: 8, 6: 60, 7: 24},
        )

        stats_sheet = workbook.create_sheet("分类统计")
        _write_table(
            stats_sheet,
            ["维度", "取值", "数量", "占比"],
            _statistics_rows(prepared, payload, threshold),
            number_formats={3: "0", 4: "0.00%"},
            widths={1: 16, 2: 32, 3: 12, 4: 12},
        )

        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_path)
        except PermissionError:
            raise ExportError(
                "FILE_LOCKED",
                "文件正被占用，无法写入：{}。请关闭 Excel 后重试。".format(output_path.name),
            )
        except OSError as exc:
            raise ExportError(
                "EXPORT_FAILED",
                "写入 Excel 失败：{}（{}）".format(output_path.name, exc),
            )
    finally:
        workbook.close()

    return output_path
