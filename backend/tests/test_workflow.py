# -*- coding: utf-8 -*-
"""新工作流：文件体检、动态类别、冲突排除与导出列名。"""

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from fastapi.testclient import TestClient

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import config
import main
from analysis import analyze_reference_file
from classifier import build_reference
from excel_io import read_reference
from export import export_results

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATA_DIR = BACKEND_DIR / "data"
REFERENCE = DATA_DIR / "分类好国标.xlsx"
PENDING = DATA_DIR / "待分类目录.xlsx"


def build_reference_workbook(path: Path, rows, header=None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "标准分类总表"
    sheet.append(header or ["标准号", "标准名称", "选用系统名称", "选用专业名称"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def build_pending_workbook(path: Path, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待分类目录"
    sheet.append(["标准号", "标准名称"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


class ReferenceAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(tempfile.mkdtemp(prefix="std_analysis_test_"))
        self.addCleanup(shutil.rmtree, self.folder, True)

    def test_real_reference_distribution_and_columns(self):
        analysis = analyze_reference_file(REFERENCE)
        self.assertEqual(analysis["columns"]["code"], "标准号")
        self.assertEqual(analysis["columns"]["name"], "标准名称")
        self.assertEqual(analysis["columns"]["system"], "选用系统名称")
        self.assertEqual(analysis["columns"]["profession"], "选用专业名称")
        self.assertEqual(analysis["raw_rows"], 105)
        self.assertEqual(analysis["training_rows"], 105)
        self.assertEqual(len(analysis["categories"]), 7)
        total = sum(category["count"] for category in analysis["categories"])
        self.assertEqual(total, 105)
        for category in analysis["categories"]:
            self.assertTrue(category["entries"])
            self.assertTrue(all(entry["code"] for entry in category["entries"]))

        same_name = analysis["quality"]["same_name_different_code"]
        self.assertEqual(len(same_name), 1)
        self.assertEqual(same_name[0]["name"], "3号喷气燃料")
        self.assertTrue(same_name[0]["same_profession"])
        self.assertEqual(same_name[0]["severity"], "info")

        same_code = analysis["quality"]["same_code_different_name"]
        self.assertEqual(len(same_code), 1)
        self.assertEqual(same_code[0]["code"], "GB/T 3131-2020")
        self.assertTrue(same_code[0]["same_profession"])
        self.assertEqual(same_code[0]["severity"], "info")

    def test_same_code_and_name_different_profession_is_excluded(self):
        path = self.folder / "冲突参照表.xlsx"
        build_reference_workbook(
            path,
            [
                ["GB/T 1-2020", "同一个名称", "系统甲", "机械装配"],
                ["GB/T 1-2020", "同一个名称", "系统乙", "表面处理"],
                ["GB/T 2-2020", "正常样本一", "系统甲", "机械装配"],
                ["GB/T 3-2020", "正常样本二", "系统乙", "表面处理"],
            ],
        )
        analysis = analyze_reference_file(path)
        self.assertEqual(analysis["quality"]["conflict_group_count"], 1)
        self.assertEqual(analysis["quality"]["conflict_row_count"], 2)
        self.assertEqual(analysis["excluded_rows"], 2)
        self.assertEqual(analysis["training_rows"], 2)
        self.assertEqual(
            {category["profession"] for category in analysis["categories"]},
            {"机械装配", "表面处理"},
        )

        records, mapping, _sheet = read_reference(path)
        reference = build_reference(records, mapping)
        self.assertEqual(reference.excluded_rows, [2, 3])
        self.assertEqual(
            sorted(reference.categories), ["机械装配", "表面处理"]
        )
        names = {name for name, _profession in reference.training_examples}
        self.assertNotIn("同一个名称", names)
        self.assertEqual(names, {"正常样本一", "正常样本二"})


class WorkflowApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)
        cls.original_output_dir = config.OUTPUT_DIR
        cls.output_dir = Path(tempfile.mkdtemp(prefix="std_workflow_test_"))
        config.OUTPUT_DIR = cls.output_dir

        reference = cls.client.post(
            "/api/reference/analyze",
            files={"reference_file": (REFERENCE.name, REFERENCE.read_bytes(), MIME)},
        )
        assert reference.status_code == 200, reference.text
        cls.reference_body = reference.json()

        pending = cls.client.post(
            "/api/pending/analyze",
            files={"pending_file": (PENDING.name, PENDING.read_bytes(), MIME)},
        )
        assert pending.status_code == 200, pending.text
        cls.pending_body = pending.json()

        response = cls.client.post(
            "/api/tasks/classify",
            json={
                "reference_id": cls.reference_body["reference_id"],
                "pending_id": cls.pending_body["pending_id"],
                "threshold": 0.7,
            },
        )
        cls.classify_status = response.status_code
        cls.classify_body = response.json()

    @classmethod
    def tearDownClass(cls):
        config.OUTPUT_DIR = cls.original_output_dir
        shutil.rmtree(cls.output_dir, ignore_errors=True)

    def test_reference_analysis_endpoint(self):
        body = self.reference_body
        self.assertTrue(body["reference_id"])
        self.assertEqual(body["raw_rows"], 105)
        self.assertEqual(len(body["categories"]), 7)
        self.assertEqual(body["columns"]["profession"], "选用专业名称")
        self.assertIn("quality", body)

    def test_pending_analysis_endpoint(self):
        body = self.pending_body
        self.assertTrue(body["pending_id"])
        self.assertEqual(body["raw_rows"], 500)
        self.assertEqual(body["unique_rows"], 179)
        self.assertEqual(body["auto_count"], 173)
        self.assertEqual(body["manual_count"], 6)
        self.assertEqual(body["multi_unit_count"], 168)
        self.assertIn("标准号 + 标准名称", body["dedup_basis"])
        self.assertTrue(body["preview"])
        manual_names = {
            item["name"] for item in body["items"] if item["status"] == "manual"
        }
        self.assertEqual(
            manual_names, {"铜及铜合金板材", "开口销", "3号喷气燃料"}
        )

    def test_classify_endpoint_returns_dynamic_categories_and_columns(self):
        self.assertEqual(self.classify_status, 200, self.classify_body)
        body = self.classify_body
        self.assertEqual(len(body["rows"]), 179)
        self.assertEqual(len(body["auto_rows"]), 173)
        self.assertEqual(len(body["manual_rows"]), 6)
        self.assertEqual(body["summary"]["raw_rows"], 500)
        self.assertEqual(body["summary"]["unique_rows"], 179)
        self.assertEqual(body["columns"]["code"], "标准号")
        self.assertTrue(all("profession" in item for item in body["categories"]))
        category_names = {
            item["profession"] for item in body["categories"]
        }
        for row in body["auto_rows"]:
            self.assertIn(row["profession"], category_names)
        for row in body["manual_rows"]:
            self.assertEqual(row["profession"], "")
            self.assertEqual(row["confidence"], 0.0)
            self.assertTrue(row["reason"])

    def test_export_appends_two_columns_to_original_pending_workbook(self):
        self.assertEqual(self.classify_status, 200, self.classify_body)
        body = self.classify_body
        overrides = {
            str(row["index"]): "金属材料" for row in body["manual_rows"]
        }
        response = self.client.post(
            "/api/export",
            json={"task_id": body["task_id"], "overrides": overrides},
        )
        self.assertEqual(response.status_code, 200, response.text)
        output_path = Path(response.json()["path"])
        workbook = openpyxl.load_workbook(output_path)
        try:
            worksheet = workbook.worksheets[0]
            headers = [cell.value for cell in worksheet[1]]
            self.assertEqual(
                headers, ["标准号", "标准名称", "未使用标注符号", "主营单位", "所属专业名称", "置信度"]
            )
            self.assertEqual(worksheet.max_row, 501)

            profession_column = len(headers) - 1
            confidence_column = len(headers)
            manual_source_rows = {
                int(row_number)
                for row in body["manual_rows"]
                for row_number in row["source_rows"]
            }
            auto_source_rows = {
                int(row_number)
                for row in body["auto_rows"]
                for row_number in row["source_rows"]
            }
            self.assertEqual(
                sum(
                    1
                    for row_number in range(2, worksheet.max_row + 1)
                    if worksheet.cell(row=row_number, column=profession_column).value
                ),
                500,
            )
            for row_number in manual_source_rows:
                self.assertEqual(
                    worksheet.cell(row=row_number, column=confidence_column).value,
                    "人工",
                )
            for row_number in list(auto_source_rows)[:5]:
                self.assertIsInstance(
                    worksheet.cell(row=row_number, column=confidence_column).value,
                    (int, float),
                )
        finally:
            workbook.close()

    def test_export_uses_uploaded_column_names(self):
        self.assertEqual(self.classify_status, 200, self.classify_body)
        folder = Path(tempfile.mkdtemp(prefix="std_export_test_"))
        self.addCleanup(shutil.rmtree, folder, True)
        reference_path = folder / "自定义表头.xlsx"
        pending_path = folder / "待分类.xlsx"
        build_reference_workbook(
            reference_path,
            [
                ["GB/T 1-2020", "不锈钢板材", "系统甲", "金属材料"],
                ["GB/T 2-2020", "铝合金棒材", "系统甲", "金属材料"],
                ["GB/T 3-2020", "阳极氧化膜", "系统乙", "表面处理"],
                ["GB/T 4-2020", "热浸镀锌层", "系统乙", "表面处理"],
            ],
            header=["标准编号", "名称", "使用系统名称", "使用专业名称"],
        )
        build_pending_workbook(pending_path, [["GB/T 1-2020", "不锈钢板材"]])

        reference = self.client.post(
            "/api/reference/analyze",
            files={"reference_file": (reference_path.name, reference_path.read_bytes(), MIME)},
        )
        pending = self.client.post(
            "/api/pending/analyze",
            files={"pending_file": (pending_path.name, pending_path.read_bytes(), MIME)},
        )
        self.assertEqual(reference.status_code, 200, reference.text)
        self.assertEqual(pending.status_code, 200, pending.text)
        reference_body = reference.json()
        self.assertEqual(reference_body["columns"]["code"], "标准编号")
        self.assertEqual(reference_body["columns"]["name"], "名称")

        response = self.client.post(
            "/api/tasks/classify",
            json={
                "reference_id": reference_body["reference_id"],
                "pending_id": pending.json()["pending_id"],
                "threshold": 0.7,
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        output = folder / "结果.xlsx"
        export_results(payload, {}, output, categories=[
            item["profession"] for item in payload["categories"]
        ])
        workbook = openpyxl.load_workbook(output)
        try:
            headers = [cell.value for cell in workbook["分类结果"][1]]
            self.assertEqual(
                headers,
                ["序号", "标准编号", "名称", "使用系统名称", "使用专业名称", "置信度", "人工复核"],
            )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
