# -*- coding: utf-8 -*-
"""FastAPI 接口与 Excel 导出的单元测试（不依赖已启动的服务器）。"""

import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from urllib.parse import quote

import openpyxl

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import config
import main
from excel_io import pick, read_reference
from export import ExportError, export_results
from fastapi.testclient import TestClient

DATA_DIR = BACKEND_DIR / "data"
REFERENCE = DATA_DIR / "分类好国标.xlsx"
PENDING = DATA_DIR / "待分类目录.xlsx"
SPREADSHEET_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

RESULT_HEADERS = [
    "序号",
    "标准号",
    "标准名称",
    "选用系统名称",
    "选用专业名称",
    "置信度",
    "人工复核",
]
SHEET_NAMES = ["分类结果", "低置信度", "数据体检", "分类统计"]


def _files():
    return {
        "pending_file": (PENDING.name, PENDING.read_bytes(), SPREADSHEET_MIME),
        "reference_file": (REFERENCE.name, REFERENCE.read_bytes(), SPREADSHEET_MIME),
    }


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.client = TestClient(main.app)
        cls.original_output_dir = config.OUTPUT_DIR
        cls.output_dir = Path(tempfile.mkdtemp(prefix="std_classify_test_output_"))
        config.OUTPUT_DIR = cls.output_dir

        cls.classify_status = None
        cls.classify_body = {}
        cls.task_id = ""
        cls.export_status = None
        cls.export_body = {}
        cls.override_index = None
        cls.override_profession = ""
        cls.original_profession = ""
        cls.reference_code = ""
        records, _mapping_rows, _sheet = read_reference(REFERENCE)
        for record in records:
            code = pick(record, "code")
            if code:
                cls.reference_code = code
                break

        response = cls.client.post(
            "/api/classify", files=_files(), data={"threshold": "0.7"}
        )
        cls.classify_status = response.status_code
        cls.classify_body = response.json()
        if response.status_code != 200:
            return

        cls.task_id = cls.classify_body.get("task_id", "")
        rows = cls.classify_body.get("rows") or []
        if not rows:
            return

        first = rows[0]
        cls.override_index = int(first["index"])
        cls.original_profession = first["profession"]
        cls.override_profession = next(
            category for category in config.CATEGORIES
            if category != cls.original_profession
        )
        export_response = cls.client.post(
            "/api/export",
            json={
                "task_id": cls.task_id,
                "overrides": {str(cls.override_index): cls.override_profession},
            },
        )
        cls.export_status = export_response.status_code
        cls.export_body = export_response.json()

    @classmethod
    def tearDownClass(cls) -> None:
        config.OUTPUT_DIR = cls.original_output_dir
        shutil.rmtree(cls.output_dir, ignore_errors=True)

    def _require_task(self) -> None:
        self.assertEqual(
            self.classify_status,
            200,
            "批量分类未成功：{}".format(self.classify_body),
        )
        self.assertTrue(self.task_id, "缺少 task_id")

    # -- 健康检查 ---------------------------------------------------------

    def test_health(self) -> None:
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["status"], "ok")
        self.assertEqual(body["version"], config.APP_VERSION)
        self.assertEqual(body["categories"], config.CATEGORIES)
        self.assertEqual(len(body["categories"]), 7)

    # -- 批量分类 ---------------------------------------------------------

    def test_classify_success(self) -> None:
        self._require_task()
        body = self.classify_body
        self.assertTrue(body["rows"])
        self.assertEqual(len(body["rows"]), body["summary"]["unique_rows"])
        self.assertAlmostEqual(body["summary"]["threshold"], 0.7, places=6)
        self.assertIn("accuracy", body["evaluation"])
        self.assertGreaterEqual(body["evaluation"]["accuracy"], 0.0)
        self.assertLessEqual(body["evaluation"]["accuracy"], 1.0)
        for key in ("new_rows", "low_rows", "conflicts", "mapping", "dedup_reports"):
            self.assertIn(key, body)
        for row in body["rows"]:
            self.assertIn(row["profession"], config.CATEGORIES)
            self.assertEqual(len(row["top3"]), 3)

    def test_classify_missing_file(self) -> None:
        response = self.client.post(
            "/api/classify",
            files={"reference_file": (REFERENCE.name, REFERENCE.read_bytes(), SPREADSHEET_MIME)},
            data={"threshold": "0.7"},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "INVALID_PARAMETER")

    def test_classify_invalid_suffix(self) -> None:
        response = self.client.post(
            "/api/classify",
            files={
                "pending_file": ("待分类.txt", b"not excel", "text/plain"),
            },
            data={"threshold": "0.7"},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "UNSUPPORTED_FILE")

    def test_classify_invalid_threshold(self) -> None:
        for value in ("0", "1", "1.5", "-0.2"):
            response = self.client.post(
                "/api/classify", files=_files(), data={"threshold": value}
            )
            self.assertEqual(response.status_code, 400, value)
            self.assertEqual(response.json()["error"]["code"], "INVALID_THRESHOLD")

    # -- 单条预测 ---------------------------------------------------------

    def test_predict_with_task(self) -> None:
        self._require_task()
        response = self.client.post(
            "/api/predict",
            json={"name": "钛合金板材", "code": "", "task_id": self.task_id},
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertIn(body["profession"], config.CATEGORIES)
        self.assertTrue(body["detail"])
        self.assertGreaterEqual(body["confidence"], 0.0)
        self.assertLessEqual(body["confidence"], 1.0)
        self.assertEqual(len(body["top3"]), 3)

    def test_predict_by_reference_code(self) -> None:
        self._require_task()
        self.assertTrue(self.reference_code, "参照表中没有可用的标准号")
        response = self.client.post(
            "/api/predict",
            json={
                "name": "完全陌生的标准名称",
                "code": self.reference_code,
                "task_id": self.task_id,
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["source"], "reference")
        self.assertIn(body["profession"], config.CATEGORIES)

    def test_predict_without_task_uses_default_reference(self) -> None:
        response = self.client.post("/api/predict", json={"name": "钛合金板材"})
        self.assertEqual(response.status_code, 200)
        self.assertIn(response.json()["profession"], config.CATEGORIES)

    def test_predict_empty_name(self) -> None:
        response = self.client.post("/api/predict", json={"name": "   "})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "INVALID_PARAMETER")

    def test_predict_unknown_task(self) -> None:
        response = self.client.post(
            "/api/predict", json={"name": "钛合金板材", "task_id": "not-exists"}
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["error"]["code"], "TASK_NOT_FOUND")

    # -- 导出 -------------------------------------------------------------

    def test_export_success(self) -> None:
        self._require_task()
        self.assertEqual(self.export_status, 200, self.export_body)
        path = Path(self.export_body["path"])
        self.assertTrue(path.is_file())
        self.assertTrue(path.name.startswith("分类结果_"))
        self.assertTrue(path.name.endswith(".xlsx"))
        self.assertEqual(
            self.export_body["download_url"],
            "/api/download/" + quote(path.name, safe=""),
        )
        self.assertEqual(self.export_body["filename"], path.name)

    def test_export_workbook_structure(self) -> None:
        self._require_task()
        self.assertEqual(self.export_status, 200, self.export_body)
        workbook = openpyxl.load_workbook(self.export_body["path"])
        try:
            self.assertEqual(workbook.sheetnames, SHEET_NAMES)

            result_sheet = workbook["分类结果"]
            headers = [cell.value for cell in result_sheet[1]]
            self.assertEqual(headers, RESULT_HEADERS)
            self.assertEqual(result_sheet.freeze_panes, "A2")
            self.assertIsNotNone(result_sheet.auto_filter.ref)

            profession_column = headers.index("选用专业名称") + 1
            manual_column = headers.index("人工复核") + 1
            confidence_column = headers.index("置信度") + 1
            matched = False
            for row in range(2, result_sheet.max_row + 1):
                if result_sheet.cell(row=row, column=1).value == self.override_index:
                    matched = True
                    self.assertEqual(
                        result_sheet.cell(row=row, column=profession_column).value,
                        self.override_profession,
                    )
                    self.assertEqual(
                        result_sheet.cell(row=row, column=manual_column).value,
                        self.override_profession,
                    )
                    self.assertEqual(
                        result_sheet.cell(row=row, column=confidence_column).number_format,
                        "0.0000",
                    )
                    break
            self.assertTrue(matched, "导出的分类结果中找不到被复核的行")

            stats_sheet = workbook["分类统计"]
            dimensions = {
                stats_sheet.cell(row=row, column=1).value
                for row in range(2, stats_sheet.max_row + 1)
            }
            self.assertIn("专业", dimensions)
            self.assertIn("系统", dimensions)
            self.assertIn("置信度区间", dimensions)

            conflict_sheet = workbook["数据体检"]
            conflict_headers = [cell.value for cell in conflict_sheet[1]]
            self.assertIn("类别", conflict_headers)
            self.assertIn("详情", conflict_headers)
        finally:
            workbook.close()

    def test_export_locked_file_raises_export_error(self) -> None:
        self._require_task()
        target = self.output_dir / "被占用的文件.xlsx"
        with patch.object(
            openpyxl.Workbook, "save", side_effect=PermissionError("locked")
        ):
            with self.assertRaises(ExportError) as context:
                export_results(self.classify_body, {}, target)
        self.assertEqual(context.exception.code, "FILE_LOCKED")
        self.assertTrue(context.exception.message)
        self.assertFalse(target.exists())

    def test_export_invalid_profession(self) -> None:
        self._require_task()
        response = self.client.post(
            "/api/export",
            json={"task_id": self.task_id, "overrides": {"1": "不存在的专业"}},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "INVALID_OVERRIDE")

    def test_export_unknown_task(self) -> None:
        response = self.client.post(
            "/api/export", json={"task_id": "not-exists", "overrides": {}}
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["error"]["code"], "TASK_NOT_FOUND")

    # -- 下载 -------------------------------------------------------------

    def test_download(self) -> None:
        self._require_task()
        self.assertEqual(self.export_status, 200, self.export_body)
        response = self.client.get(self.export_body["download_url"])
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn(
            "spreadsheetml", response.headers.get("content-type", "")
        )

    def test_download_missing_file(self) -> None:
        response = self.client.get(
            "/api/download/" + quote("不存在的文件.xlsx", safe="")
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["error"]["code"], "FILE_NOT_FOUND")

    def test_download_rejects_traversal(self) -> None:
        for name in ("../main.py", "..\\main.py", "sub/../../main.py"):
            response = self.client.get("/api/download/" + quote(name, safe=""))
            self.assertIn(response.status_code, (400, 404), name)
            self.assertIn("error", response.json(), name)


if __name__ == "__main__":
    unittest.main()
