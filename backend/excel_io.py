# -*- coding: utf-8 -*-
"""Excel 读取、列识别与待分类目录去重。"""

import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import openpyxl

from config import COLUMN_ALIASES, MAPPING_SHEET, REFERENCE_SHEET
from text import normalize_code, normalize_name, to_text

# GB / GB/T、GJB / GJB/T 只差写法，视为同一标准。
_VARIANT_PREFIX_RE = re.compile(r"^(GJB|GB)T(?=\d)")


def _variant_code_key(code: Any) -> str:
    """把 GB/T、GJB/T 前缀折叠成 GB、GJB，用于识别写法变体。"""
    return _VARIANT_PREFIX_RE.sub(r"\1", normalize_code(code))


def variant_code_key(code: Any) -> str:
    """公开的写法变体标准号键。"""
    return _variant_code_key(code)


class ExcelError(Exception):
    """可预期的 Excel 读取错误，携带稳定错误码。"""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _all_aliases() -> Tuple[str, ...]:
    aliases: List[str] = []
    for values in COLUMN_ALIASES.values():
        aliases.extend(values)
    return tuple(aliases)


def _header_matches(header: Sequence[str], field: str) -> bool:
    return any(alias in header for alias in COLUMN_ALIASES[field])


def _required_fields_error(source: str, header: Sequence[str], fields: Sequence[str]) -> None:
    missing: List[str] = []
    for field in fields:
        if not _header_matches(header, field):
            missing.append(COLUMN_ALIASES[field][0])
    if missing:
        raise ExcelError(
            "INVALID_FILE",
            "{}缺少必需的列：{}".format(source, "、".join(missing)),
        )


def _unique_header(values: Sequence[Any]) -> List[str]:
    header: List[str] = []
    seen: Counter = Counter()
    for value in values:
        name = to_text(value)
        if not name:
            header.append("")
            continue
        seen[name] += 1
        header.append(name if seen[name] == 1 else "{}_{}".format(name, seen[name]))
    return header


def read_table(
    path: Path,
    sheet_name: Optional[str] = None,
    required_fields: Sequence[str] = (),
    source: str = "工作表",
) -> Tuple[List[str], List[Dict[str, Any]], str]:
    """读取工作表，返回 (表头, 记录, 工作表名)。

    表头不要求在第一行：会扫描前 20 行，找第一行包含已知列名的行。
    """
    file_path = Path(path)
    if not file_path.exists():
        raise ExcelError("FILE_NOT_FOUND", "文件不存在：{}".format(file_path))
    if file_path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelError("INVALID_FILE", "只支持 .xlsx / .xlsm 文件：{}".format(file_path.name))

    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - 统一转成可展示的中文错误
        raise ExcelError("INVALID_FILE", "无法读取 Excel 文件：{}（{}）".format(file_path.name, exc))

    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise ExcelError("SHEET_NOT_FOUND", "工作表不存在：{}".format(sheet_name))
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]

        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return [], [], worksheet.title

        aliases = set(_all_aliases())
        header_index = None
        for index, row in enumerate(rows[:20]):
            texts = {to_text(cell) for cell in row}
            if texts & aliases:
                header_index = index
                break
        if header_index is None:
            # 没有已知列名时仍按第一行当表头，让后续校验给出明确缺列错误。
            for index, row in enumerate(rows[:20]):
                if any(to_text(cell) for cell in row):
                    header_index = index
                    break
        if header_index is None:
            return [], [], worksheet.title

        header = _unique_header(rows[header_index])
        _required_fields_error(source, header, required_fields)

        records: List[Dict[str, Any]] = []
        for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if all(cell is None or to_text(cell) == "" for cell in row):
                continue
            record: Dict[str, Any] = {"_row": row_index}
            for column_index, name in enumerate(header):
                if not name:
                    continue
                record[name] = row[column_index] if column_index < len(row) else None
            records.append(record)
        if records:
            # 记录表头所在的 Excel 行号，导出时用于在原表追加列。
            records[0]["_header_row"] = header_index + 1
        return header, records, worksheet.title
    finally:
        workbook.close()


def pick(record: Dict[str, Any], field: str) -> str:
    """按列名别名取第一个非空值。"""
    for alias in COLUMN_ALIASES[field]:
        if alias in record:
            value = to_text(record.get(alias))
            if value:
                return value
    return ""


def _merge_writing_variants(
    items: List[Dict[str, Any]],
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """合并名称、年份和编号相同、仅 GB/GB/T（GJB/GJB/T）写法不同的条目。

    保留首次出现的标准号，合并重复次数与来源行。
    """
    merged: List[Dict[str, Any]] = []
    canonical_by_key: Dict[Tuple[str, str], Dict[str, Any]] = {}
    reports: List[Dict[str, Any]] = []

    for item in items:
        code = item["code"]
        variant_key = _variant_code_key(code) if code else ""
        if not variant_key:
            merged.append(item)
            continue

        group_key = (normalize_name(item["name"]), variant_key)
        canonical = canonical_by_key.get(group_key)
        if canonical is None:
            canonical_by_key[group_key] = item
            item["_variant_members"] = [{
                "code": code,
                "count": item["count"],
                "rows": list(item["source_rows"]),
            }]
            merged.append(item)
            continue

        canonical["count"] += item["count"]
        canonical["source_rows"].extend(item["source_rows"])
        canonical["codes"].update(item["codes"])
        canonical["_variant_members"].append({
            "code": code,
            "count": item["count"],
            "rows": list(item["source_rows"]),
        })

    for item in merged:
        members = item.pop("_variant_members", [])
        if len(members) <= 1:
            continue
        detail = "；".join(
            "{}（{} 条，Excel 行：{}）".format(
                member["code"],
                member["count"],
                "、".join(str(row) for row in member["rows"]) or "未知",
            )
            for member in members
        )
        reports.append({
            "type": "待分类-写法变体合并",
            "code": item["code"],
            "name": item["name"],
            "count": item["count"],
            "detail": detail,
            "action": "合并为 1 条，保留首条标准号 {}".format(item["code"]),
        })

    return merged, reports


def deduplicate(
    records: Sequence[Dict[str, Any]],
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """按标准号去重；无标准号时按名称去重。返回 (条目, 报告)。"""
    items: List[Dict[str, Any]] = []
    index_by_key: Dict[str, Dict[str, Any]] = {}

    for record in records:
        code = pick(record, "code")
        name = pick(record, "name")
        if not name:
            continue

        code_key = normalize_code(code)
        key = code_key if code_key else "NAME:" + normalize_name(name)
        row_number = record.get("_row", "")

        if key not in index_by_key:
            item = {
                "code": code,
                "name": name,
                "count": 0,
                "row_number": row_number,
                "source_rows": [],
                "codes": set(),
            }
            index_by_key[key] = item
            items.append(item)

        item = index_by_key[key]
        item["count"] += 1
        if row_number != "":
            item["source_rows"].append(row_number)
        if code:
            item["codes"].add(normalize_code(code))

    reports: List[Dict[str, Any]] = []
    for item in items:
        if item["count"] > 1:
            reports.append({
                "type": "待分类-重复记录",
                "code": item["code"] or "无标准号",
                "name": item["name"],
                "count": item["count"],
                "detail": "Excel 行：{}".format(
                    "、".join(str(row) for row in item["source_rows"]) or "未知"
                ),
                "action": "合并为 1 条",
            })

    items, variant_reports = _merge_writing_variants(items)
    reports.extend(variant_reports)

    # 同名不同标准号：写法变体已合并，同一标准号的不同写法不再重复报告。
    codes_by_name: Dict[str, Dict[str, str]] = defaultdict(dict)
    for item in items:
        code = item["code"]
        if not code:
            continue
        variant_key = _variant_code_key(code)
        codes_by_name[normalize_name(item["name"])].setdefault(variant_key, normalize_code(code))

    for name_key, code_map in codes_by_name.items():
        if len(code_map) > 1:
            reports.append({
                "type": "待分类-同名不同标准号",
                "code": "、".join(sorted(code_map.values())),
                "name": name_key,
                "count": len(code_map),
                "detail": "同一名称对应多个标准号",
                "action": "保留为不同标准，不合并",
            })

    for item in items:
        item.pop("codes", None)

    return items, reports


def _row_sort_key(value: Any) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def group_pending_records(
    records: Sequence[Dict[str, Any]],
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Dict[str, Any]]:
    """按「标准号 + 标准名称」归并待分类数据，并标记可自动/需人工。

    - 同一标准号 + 同一标准名称（GB/T 与 GB 写法视为同一号）归为一组；
    - 同名不同号、同号不同名、缺少标准号/名称的组标记为 manual；
    - 同号同名但主营单位不同的组仍归为同一标准，标记 needs_confirmation；
    - 每个原始行都保留在 source_rows 中，供导出时逐行回填。
    """
    rows: List[Dict[str, Any]] = []
    for record in records:
        code = pick(record, "code")
        name = pick(record, "name")
        unit = pick(record, "unit")
        if not (code or name or unit):
            continue
        rows.append({
            "code": code,
            "name": name,
            "unit": unit,
            "row": record.get("_row", ""),
            "code_family": _variant_code_key(code) if code else "",
            "name_key": normalize_name(name),
        })

    groups: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for row in rows:
        key = (
            row["code_family"] or "NAME:" + row["name_key"],
            row["name_key"],
        )
        group = groups.get(key)
        if group is None:
            group = {
                "code": row["code"],
                "name": row["name"],
                "code_family": row["code_family"],
                "name_key": row["name_key"],
                "unit_counts": Counter(),
                "source_rows": [],
                "count": 0,
                "first_row": row["row"],
            }
            groups[key] = group
        group["count"] += 1
        if row["row"] != "":
            group["source_rows"].append(row["row"])
            if group["first_row"] == "":
                group["first_row"] = row["row"]
        if row["unit"]:
            group["unit_counts"][row["unit"]] += 1

    names_by_code: Dict[str, set] = defaultdict(set)
    codes_by_name: Dict[str, set] = defaultdict(set)
    for row in rows:
        if row["code_family"]:
            names_by_code[row["code_family"]].add(row["name_key"])
            codes_by_name[row["name_key"]].add(row["code_family"])

    items: List[Dict[str, Any]] = []
    for group in groups.values():
        reasons: List[str] = []
        if not group["name"]:
            reasons.append("缺少标准名称")
        if not group["code"]:
            reasons.append("缺少标准号")
        if (
            group["code_family"]
            and len(names_by_code[group["code_family"]]) > 1
        ):
            reasons.append("同标准号不同名称")
        if group["name_key"] and len(codes_by_name[group["name_key"]]) > 1:
            reasons.append("同名不同标准号")

        units = [unit for unit, _count in group["unit_counts"].most_common()]
        status = "manual" if reasons else "auto"
        items.append({
            "code": group["code"],
            "name": group["name"],
            "units": units,
            "unit_counts": dict(group["unit_counts"]),
            "count": group["count"],
            "source_rows": sorted(group["source_rows"], key=_row_sort_key),
            "status": status,
            "reason": "；".join(reasons),
            "needs_confirmation": status == "auto" and len(units) > 1,
            "first_row": group["first_row"],
        })

    items.sort(key=lambda item: _row_sort_key(item["first_row"]))
    for index, item in enumerate(items, start=1):
        item["index"] = index

    reports: List[Dict[str, Any]] = []
    for item in items:
        distinct_units = len(item["unit_counts"]) or 1
        duplicate_rows = item["count"] - distinct_units
        if duplicate_rows > 0:
            reports.append({
                "type": "待分类-完全重复行",
                "code": item["code"] or "无标准号",
                "name": item["name"],
                "count": duplicate_rows,
                "detail": "同一标准号、名称、主营单位重复出现",
                "action": "合并为同一标准，导出时逐行回填",
            })
        if item["status"] == "manual":
            reports.append({
                "type": "待分类-需人工分类",
                "code": item["code"] or "无标准号",
                "name": item["name"] or "无名称",
                "count": item["count"],
                "detail": item["reason"],
                "action": "不进入自动分类，由用户选择所属专业名称",
            })

    exact_duplicate_rows = sum(
        item["count"] - (len(item["unit_counts"]) or 1)
        for item in items
    )
    auto_count = sum(1 for item in items if item["status"] == "auto")
    manual_count = len(items) - auto_count
    multi_unit_count = sum(1 for item in items if item["needs_confirmation"])
    meta = {
        "raw_rows": len(rows),
        "unique_rows": len(items),
        "group_count": len(items),
        "auto_count": auto_count,
        "manual_count": manual_count,
        "multi_unit_count": multi_unit_count,
        "exact_duplicate_rows": exact_duplicate_rows,
        "dedup_basis": (
            "按「标准号 + 标准名称」归并；GB/T 与 GB 写法视为同一标准号；"
            "不合并同名不同号、同号不同名；主营单位不同的行保留在同一标准下，并标记为待确认。"
        ),
    }
    return items, reports, meta


def load_pending(path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """读取待分类目录并去重。"""
    _header, records, _sheet = read_table(
        path, required_fields=("name",), source="待分类目录"
    )
    items, reports = deduplicate(records)
    if not items:
        raise ExcelError("INVALID_FILE", "待分类目录没有可用的标准名称")
    return items, reports


def read_reference_full(
    path: Path,
) -> Tuple[List[str], List[Dict[str, Any]], List[Dict[str, Any]], str]:
    """读取参照表，返回 (表头, 总表记录, 映射表记录, 工作表名)。"""
    try:
        header, records, sheet = read_table(
            path, REFERENCE_SHEET, required_fields=("name", "profession"), source="参照表"
        )
    except ExcelError as exc:
        if exc.code != "SHEET_NOT_FOUND":
            raise
        header, records, sheet = read_table(
            path, required_fields=("name", "profession"), source="参照表"
        )

    mapping_rows: List[Dict[str, Any]] = []
    try:
        _mapping_header, mapping_rows, _mapping_sheet = read_table(
            path, MAPPING_SHEET, required_fields=("profession", "system"), source="专业与系统映射"
        )
    except ExcelError:
        mapping_rows = []

    return header, records, mapping_rows, sheet


def read_reference(
    path: Path,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], str]:
    """读取参照表总表与专业/系统映射表。"""
    _header, records, mapping_rows, sheet = read_reference_full(path)
    return records, mapping_rows, sheet


def count_raw_rows(records: Sequence[Dict[str, Any]]) -> int:
    """原始数据行数（含重复）。"""
    return len(records)
